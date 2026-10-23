# Builds "Tithi Family Notebook.xlsx": the shared sheet the WhatsApp job reads.
# Upload it to Google Drive and open it with Google Sheets; the dropdowns carry over.
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

lists = json.load(open(sys.argv[1], encoding='utf-8'))
out = sys.argv[2]

PAKSHAS = ['Shukla · శుక్ల', 'Bahula (Krishna) · బహుళ']
OCCASIONS = ['Birthday · పుట్టినరోజు', 'Wedding anniversary · పెళ్లిరోజు', 'Aabdikam · ఆబ్దికం', 'Other · ఇతర']
KINDS = ['Festival', 'Telugu date every year', 'Tithi every month']
ROWS = 300

wb = Workbook()
head_font = Font(bold=True, color='FFFFFF')
head_fill = PatternFill('solid', fgColor='232B52')
hint_font = Font(italic=True, color='5B6283')

def header(ws, cols, widths, hint):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    ws.cell(row=1, column=len(cols) + 2, value=hint).font = hint_font

# Lists tab (dropdown sources)
ls = wb.active
ls.title = 'Lists'
columns = {'A': ('Masam', lists['masas']), 'B': ('Paksham', PAKSHAS), 'C': ('Tithi', lists['tithis']),
           'D': ('Festival', [f['label'] for f in lists['fest']]), 'E': ('Occasion', OCCASIONS), 'F': ('Comes back on', KINDS)}
for col, (title, values) in columns.items():
    ls[f'{col}1'] = title
    ls[f'{col}1'].font = Font(bold=True)
    for i, v in enumerate(values, 2):
        ls[f'{col}{i}'] = v
    ls.column_dimensions[col].width = 34
ref = lambda col, n: f'=Lists!${col}$2:${col}${n + 1}'

def dropdown(ws, col, src):
    dv = DataValidation(type='list', formula1=src, allow_blank=True, showErrorMessage=True,
                        errorTitle='Pick from the list', error='Choose a value from the dropdown.')
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}{ROWS}')

# Family tab
fam = wb.create_sheet('Family', 0)
header(fam, ['Name', 'Occasion', 'Masam', 'Paksham', 'Tithi', 'Extra things to do (optional)'],
       [22, 30, 26, 24, 34, 44],
       'One row per person. Pick masam, paksham and tithi of the birth / wedding / passing. Comes back every year.')
dropdown(fam, 'B', ref('E', len(OCCASIONS)))
dropdown(fam, 'C', ref('A', 12))
dropdown(fam, 'D', ref('B', 2))
dropdown(fam, 'E', ref('C', 15))

# Notes tab
nt = wb.create_sheet('Notes', 1)
header(nt, ['Comes back on', 'Festival (if festival)', 'Masam (if Telugu date)', 'Paksham', 'Tithi', 'What to do', 'Added by'],
       [24, 40, 26, 24, 34, 44, 14],
       'Festival: pick the festival. Telugu date: pick masam + paksham + tithi. Every month: pick paksham + tithi.')
dropdown(nt, 'A', ref('F', 3))
dropdown(nt, 'B', ref('D', len(lists['fest'])))
dropdown(nt, 'C', ref('A', 12))
dropdown(nt, 'D', ref('B', 2))
dropdown(nt, 'E', ref('C', 15))
# Already-saved note, moved over from the Tithi Reminders page
mukkoti = next(f['label'] for f in lists['fest'] if f['id'] == 'mukkoti')
nt.append(['Festival', mukkoti, '', '', '', 'Rama Japam', ''])

ls.sheet_state = 'hidden'
wb.save(out)
print('saved', out)
